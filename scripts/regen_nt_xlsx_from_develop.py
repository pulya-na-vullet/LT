#!/usr/bin/env python3
"""Regenerate NT Excel from ump-ncins-pa develop (workers + IT + BPMN)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path("/workspace/nt-prerequisites-xlsx")
DATE = "UMP · 21 июля 2026 · код develop ump-ncins-pa@4a921a1653e + IT stubs"

HEADER_FILL = PatternFill("solid", fgColor="D6EAF8")
TITLE_FONT = Font(name="Calibri", size=14, bold=True)
SUB_FONT = Font(name="Calibri", size=11, italic=True)
SECTION_FONT = Font(name="Calibri", size=12, bold=True)
HEADER_FONT = Font(name="Calibri", size=11, bold=True)
CELL_FONT = Font(name="Calibri", size=11)
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
COLS = ["#", "Пререквизит", "Значение", "Пример", "Где взять", "Предпочтительная роль заполняющего"]
COL_WIDTHS = {"A": 5, "B": 32, "C": 58, "D": 45, "E": 28, "F": 18}

WHERE_ALGO = "Ссылка на алгоритм работы метода в git или confluence."
WHERE_SLA = "Получить от заказчика доработки. Если заказчик неизвестен, определяем экспертно."
WHERE_CHPN = "Аналитику нужно определить: 1) макс вызовов/час; 2) среднее вызовов/час; 3) способ получения оценки."
WHERE_CHPN6 = "Само НТ проводится по профилю прогнозируемой через полгода ЧПН."
WHERE_DEP = "Может совпадать с ЧПН."
WHERE_DEP6 = "Может совпадать с прогнозируемой через полгода ЧПН."
WHERE_CALL = "Сначала нужно добыть живой пример вызова сервиса (локально/стенд) и параметризовать."
WHERE_RESP = "Нужны примеры ответов, которые можно считать ожидаемыми."
WHERE_MUT = "Перечислить поля/таблицы/ресурсы, которые создаются/перезаписываются (кроме логирования)."
WHERE_MIG = "Опционально. Скрипт готовит разработчик только однажды перед боем."
WHERE_STUB = "Для зависимостей, которые нужно/можно глушить — прикрепить моки."
WHERE_LAT = "Потребуется для имитации нагрузки на внешние зависимости."
WHERE_SCEN = "Опционально. Готовит аналитик, если тестируется бизнес-процесс, а не метод."


def style_header_row(ws, row: int, ncols: int = 6) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = THIN


def set_widths(ws) -> None:
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width


def write_title(ws, title: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = DATE
    ws["A2"].font = SUB_FONT


def write_toc(ws, title: str, purpose: str) -> None:
    write_title(ws, title)
    ws["A4"] = "Table of Contents"
    ws["A4"].font = SECTION_FONT
    ws["A5"] = "1 Согласование"
    ws["A6"] = "2 1. Минимальная информация для определения необходимости проведения НТ"
    ws["A7"] = (
        "3 2. Полная информация для проведения НТ "
        "(заполняется, только если получено заключение от инженера о необходимости проведения НТ)"
    )
    ws["A9"] = purpose
    ws["A9"].font = CELL_FONT
    ws["A9"].alignment = WRAP
    set_widths(ws)
    ws.row_dimensions[9].height = 55


def write_agreement(ws, title: str) -> None:
    write_title(ws, title)
    ws["A4"] = "1 Согласование"
    ws["A4"].font = SECTION_FONT
    ws["A5"] = "Аналитик"
    ws["B5"] = "Согласование"
    style_header_row(ws, 5, 2)
    ws["A6"] = "TBD"
    ws["B6"] = "TBD"
    for col in range(1, 3):
        ws.cell(row=6, column=col).font = CELL_FONT
        ws.cell(row=6, column=col).border = THIN
    set_widths(ws)


def write_table(ws, title: str, section: str, rows: list[tuple]) -> None:
    write_title(ws, title)
    ws.cell(row=4, column=1, value=section).font = SECTION_FONT
    for i, h in enumerate(COLS, 1):
        ws.cell(row=5, column=i, value=h)
    style_header_row(ws, 5)
    for r_i, row in enumerate(rows, 6):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
        text = str(row[2]) if len(row) > 2 else ""
        ws.row_dimensions[r_i].height = min(240, max(70, 14 + text.count("\n") * 11))
    set_widths(ws)


def write_bpmn_sheet(ws, title: str, section: str, rows: list[tuple], headers=None) -> None:
    write_title(ws, title)
    ws.cell(row=4, column=1, value=section).font = SECTION_FONT
    headers = headers or [
        "#",
        "BPMN / JobWorker",
        "Name",
        "Type / HTTP",
        "Request / Input",
        "Response / Output",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=5, column=i, value=h)
    style_header_row(ws, 5)
    for r_i, row in enumerate(rows, 6):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
        ws.row_dimensions[r_i].height = 70
    set_widths(ws)
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 36
    ws.column_dimensions["E"].width = 42
    ws.column_dimensions["F"].width = 42


def build(meta: dict) -> Workbook:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Содержание"
    write_toc(ws0, meta["title"], meta["purpose"])
    write_agreement(wb.create_sheet("1. Согласование"), meta["title"])
    write_table(
        wb.create_sheet("2. Минимальная информация"),
        meta["title"],
        "1. Минимальная информация для определения необходимости проведения НТ",
        meta["min_rows"],
    )
    write_table(
        wb.create_sheet("3. Полная информация"),
        meta["title"],
        "2. Полная информация для проведения НТ (заполняется, только если получено заключение от инженера о необходимости проведения НТ)",
        meta["full_rows"],
    )
    write_bpmn_sheet(
        wb.create_sheet("Методы и примеры (код)"),
        meta["title"],
        meta["methods_section"],
        meta["method_rows"],
    )
    return wb


LOAD = (
    "База нагрузки СА: 15000 заявок/мес SFA → ≈89/час среднее, пик ≈180/час. "
    "Timeout-ветки ~5% (из них ~95% signing)."
)

# ---------------------------------------------------------------------------
PROCESSES: list[dict] = []

# MAIN — BPMN отсутствует в develop; оркестрация по дочерним + ответы СА
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-main-ma-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-main-ma-ncins-pa",
        "purpose": (
            "Назначение: оркестратор мультизаявки ncins. "
            "В develop ump-ncins-pa BPMN ump-main-ma-ncins-pa.bpmn ОТСУТСТВУЕТ "
            "(есть только 5 дочерних процессов). Ниже — по дочерним BPMN/воркерам + СА."
        ),
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "[ump-main-ma-ncins-pa] Оркестратор (BPMN файла нет в develop).\n"
                    "По дочерним процессам в репе: prepare → signing → payment → finalisation; "
                    "на timeout — delete.\n"
                    "Дочерние job types реализованы в ump-ncins-pa (*Worker.java)."
                ),
                "Нужен BPMN main-ma из отдельной ветки/PR",
                WHERE_ALGO,
                "Разработчик",
            ),
            (
                "2",
                "SLA по времени отклика",
                "E2E TBD с НТ. Sync дочерних: несколько секунд. В develop payment timer = PT15M (BPMN).",
                "TBD",
                WHERE_SLA,
                "Аналитик",
            ),
            ("3", "Частота планируемой нагрузки (ЧПН)", LOAD, "89 / 180 в час", WHERE_CHPN, "Аналитик"),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                (
                    "prepare/signing ~89–180; payment/finalisation ~85–170; delete ~4–5.\n"
                    "HTTP из дочерних: GET /applications/{id}; PUT /products/{id}; POST /v1/ins-contracts; "
                    "внешние connectors (EA, transfer-control)."
                ),
                "См. листы дочерних процессов",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "HTTP нет. Zeebe create instance bpmnProcessId=ump-main-ma-ncins-pa.\n"
                    "Vars: businessKey (UUID), productCode=NON_CREDIT_INSURANCE.\n"
                    "ВНИМАНИЕ: BPMN main в develop-архиве нет — пример старта уточнить по ветке с main-ma."
                ),
                '{"businessKey":"<uuid>","productCode":"NON_CREDIT_INSURANCE"}',
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                "Happy-path: result=SUCCESS. Reject: BANK_REJECT (по СА/предыдущему PR). В develop коде main нет.",
                "SUCCESS / BANK_REJECT",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "Через дочерние: PUT products (policyLink), POST ins-contracts; стадии/Kafka — вне этого сервиса.",
                "policyLink / договор АБ",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "Payment hold/create — stub в PaymentService. Delete AC — stub. EA/transfer — внешние connectors. Application facade / products / NIB — реальные Feign.",
                "См. дочерние Excel",
                WHERE_STUB,
                "Разработчик",
            ),
            (
                "11",
                "Ожидаемое время отклика замоканных внешних зависимостей",
                "Sync несколько секунд; async message на НТ сократить (≤~3 сек signing — по СА).",
                "несколько сек",
                WHERE_LAT,
                "Разработчик",
            ),
            (
                "12",
                "Сценарий тестирования",
                "E2E по дочерним: prepare→signing→payment→finalisation; timeout→delete. Без BPMN main — тестировать дочерние по отдельности или подложить main из другой ветки.",
                "5 дочерних BPMN в develop",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "methods_section": "Дочерние JobWorker в ump-ncins-pa (main BPMN отсутствует)",
        "method_rows": [
            ("1", "get-application-data", "ApplicationDataWorker", "GET /applications/{id}", "businessKey, processType", "getProcessParams + поля по типу"),
            ("2", "...get-report-data", "PrepareDataForPrintFormWorker", "локально XML→Base64", "поля ПФ", "documents[].reportData"),
            ("3", "update-product", "ProductDataUpdateWorker", "PUT /products/{id}", "productId, acDocuments, productInfo", "getProcessParams"),
            ("4", "...set-hold / create-payment", "InsurancePaymentWorker", "STUB (нет HTTP)", "Map variables", '{"code":"SUCCESS"}'),
            ("5", "...create-contract", "FinalisationWorker", "POST /v1/ins-contracts", "ContractCreate vars", '{"result":"SUCCESS"}'),
            ("6", "...delete-documents", "DeleteDocumentsWorker", "STUB AC 1.0", "document.documentLink", '{"documentResponse":{"status":"1"}}'),
            ("7", "connectors", "внешние", "ea-send / transfer-control", "acId / messageName", "вне ump-ncins-pa"),
        ],
    }
)

# PREPARE
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-prepare-documents-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-prepare-documents-ncins-pa",
        "purpose": "BPMN ump-prepare-documents-ncins-pa + воркеры ApplicationData / PrepareDataForPrintForm / ProductDataUpdate + Call Activity generate-and-save.",
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "BPMN develop:\n"
                    "1) get-application-data (processType=PREPARE_DOCUMENTS) → ApplicationDataWorker\n"
                    "2) ump-prepare-documents-ncins-pa.get-report-data → PrepareDataForPrintFormWorker\n"
                    "3) Call Activity ump-generate-and-save-document-pa (внешний процесс)\n"
                    "4) update-product → ProductDataUpdateWorker (policyLink из acDocuments[0].acId)"
                ),
                "src/main/resources/bpmn/ump-prepare-documents-ncins-pa.bpmn",
                WHERE_ALGO,
                "Разработчик",
            ),
            ("2", "SLA по времени отклика", "Все шаги sync (кроме внешнего Call Activity). Несколько секунд/вызов (TBD).", "несколько сек", WHERE_SLA, "Аналитик"),
            ("3", "Частота планируемой нагрузки (ЧПН)", LOAD + " 1 документ/заявка.", "89 / 180", WHERE_CHPN, "Аналитик"),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                (
                    "[GET ump-application-facade /applications/{id}] ~89–180\n"
                    "[get-report-data локально] ~89–180\n"
                    "[Call Activity generate-and-save] ~89–180\n"
                    "[PUT ump-products /products/{id}] ~89–180"
                ),
                "applicationFacade.url / products.url",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "Старт процесса / Call Activity из main.\n"
                    "Шаг1 Camunda vars: {\"businessKey\":\"550e8400-e29b-41d4-a716-446655110003\",\"processType\":\"PREPARE_DOCUMENTS\"}\n"
                    "HTTP ApplicationDataWorker:\n"
                    "GET {applicationFacade.url}/applications/{businessKey}?include=PARTICIPANT&include=PRODUCT\n"
                    "url prod-like: http://ump-application-facade.ump.svc.cluster.local\n"
                    "Шаг2 get-report-data input (из IT): programId, fullName, inn, email, contractNumber, dates, sums, paymentAccount.\n"
                    "Шаг4 update-product: productId, acDocuments[{type,acId}], productInfo."
                ),
                "IT: ApplicationDataWorkerIT, PrepareDataForPrintFormWorkerIT, ProductDataUpdateWorkerIT",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                (
                    "get-application-data: getProcessParams=SUCCESS + поля PREPARE_DOCUMENTS "
                    "(productId, fullName, inn, contractNumber, …). Stub: stubs/ump/ump_application_response_200.json\n"
                    "get-report-data out:\n"
                    '{"serviceId":"<businessKey>","serviceCode":"APPLICATION","productCode":"NON_CREDIT_INSURANCE",'
                    '"documents":[{"documentType":"CONTRACT_ACCOUNT_BLOCK","reportData":"<Base64 XML>","isWriteInEa":false}]}\n'
                    "update-product: {\"getProcessParams\":\"SUCCESS\"} (ERROR при exception).\n"
                    "Ошибки facade → getProcessParams=ERROR."
                ),
                "SUCCESS / ERROR; XML в Base64",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "PUT /products/{id}: в productProperties.policyLink пишется acDocuments[0].acId (код ProductDataUpdateWorker).",
                "policyLink = acId",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "Application facade и products — реальные Feign (на НТ WireMock/стенд). Call Activity generate-and-save — внешний процесс (ПФ+AC), мок/стенд уточнять.",
                "WireMock как в IT",
                WHERE_STUB,
                "Разработчик",
            ),
            ("11", "Ожидаемое время отклика замоканных внешних зависимостей", "Несколько секунд на sync-вызов (TBD).", "несколько сек", WHERE_LAT, "Разработчик"),
            (
                "12",
                "Сценарий тестирования",
                "1) Success: get→report→generate-save→update.\n2) facade error → ERROR.\n3) putProduct exception → ERROR.\n4) пик ~180/час.",
                "IT покрывают шаги 1/2/4 частично",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "methods_section": "Методы prepare — JobWorker + HTTP (develop + IT)",
        "method_rows": [
            (
                "1",
                "get-application-data",
                "ApplicationDataWorker",
                "GET /applications/{id}?include=PARTICIPANT&include=PRODUCT",
                '{"businessKey":"550e8400-...","processType":"PREPARE_DOCUMENTS"}',
                "ApplicationResponseDto → getProcessParams=SUCCESS + product fields",
            ),
            (
                "2",
                "ump-prepare-documents-ncins-pa.get-report-data",
                "PrepareDataForPrintFormWorker",
                "нет HTTP (XML+Base64)",
                "fullName/inn/email/contractNumber/sums/paymentAccount (IT)",
                "documents[0].documentType=CONTRACT_ACCOUNT_BLOCK, reportData=Base64",
            ),
            (
                "3",
                "Call Activity",
                "ump-generate-and-save-document-pa",
                "внешний процесс",
                "serviceId=businessKey, serviceCode=APPLICATION, documents",
                "acDocuments, acRequestId, docsResult",
            ),
            (
                "4",
                "update-product",
                "ProductDataUpdateWorker",
                "PUT /products/{id}",
                '{"productId":"660e...","acDocuments":[{"type":"CONTRACT_ACCOUNT_BLOCK","acId":"770e..."}],"productInfo":{}}',
                '{"getProcessParams":"SUCCESS"}',
            ),
        ],
    }
)

# SIGNING
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-signing-documents-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-signing-documents-ncins-pa",
        "purpose": (
            "BPMN ump-signing-documents-ncins-pa (develop): get-application-data + "
            "connector transfer-control + message SigningDocs. В этом сервисе только ApplicationDataWorker."
        ),
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "develop BPMN:\n"
                    "1) get-application-data (SIGNING) → ApplicationDataWorker\n"
                    "2) XOR: если systemCode=\"AlfaOffice\" → End; иначе\n"
                    "3) connector ump-process-transfer-connector.transfer-control.v1 "
                    "(messageName=SigningDocs, useProcessInstanceKey=true, systemCode=productCode, stopInIncident=true)\n"
                    "4) receiveTask SigningDocs, correlationKey=processInstanceKey\n"
                    "В develop НЕТ update-product и НЕТ timer PT25M в этом BPMN (отличие от позднего PR)."
                ),
                "ump-signing-documents-ncins-pa.bpmn",
                WHERE_ALGO,
                "Разработчик",
            ),
            ("2", "SLA по времени отклика", "Sync get — несколько сек. Async wait SigningDocs — в проде до минут; на НТ сократить (≤~3 сек по СА).", "≤~3 сек на НТ", WHERE_SLA, "Аналитик"),
            ("3", "Частота планируемой нагрузки (ЧПН)", LOAD, "89 / 180", WHERE_CHPN, "Аналитик"),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                "[get-application-data / GET applications] ~89–180\n[transfer-control connector / Kafka] ~89–180\n[Receive SigningDocs] ~89–180",
                "внешний connector не в ump-ncins-pa",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "Camunda: businessKey + processType=SIGNING.\n"
                    "HTTP: GET /applications/{id}?include=PARTICIPANT&include=PRODUCT\n"
                    "Output SIGNING (маппер): processInstanceKey + getProcessParams.\n"
                    "Далее внешний connector + message SigningDocs (corr=processInstanceKey в develop BPMN).\n"
                    "Полный JSON Kafka — TBD (контракт аналитики / connector)."
                ),
                "ApplicationDataWorkerIT (processType можно SIGNING)",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                "Worker: getProcessParams=SUCCESS|ERROR. После message — End. AlfaOffice — короткий путь без connector.",
                "SUCCESS / ERROR",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "В ump-ncins-pa на signing update-product нет. Изменения стадии — через внешний connector/Kafka (вне репы).",
                "UMP DB через этот сервис не меняется",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "Facade — стенд/WireMock. transfer-control + SigningDocs — подкладывать message в Zeebe/Kafka на НТ (JSON TBD).",
                "corrKey=processInstanceKey (develop)",
                WHERE_STUB,
                "Разработчик",
            ),
            ("11", "Ожидаемое время отклика замоканных внешних зависимостей", "get sync несколько сек; message на НТ ≤~3 сек.", "≤3 сек", WHERE_LAT, "Разработчик"),
            (
                "12",
                "Сценарий тестирования",
                "1) Success: get→transfer→подложить SigningDocs.\n2) AlfaOffice shortcut.\n3) Без message — зависание (timer в develop BPMN нет).",
                "Сверить с PR, если добавят timer/update-product",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "methods_section": "Методы signing (develop)",
        "method_rows": [
            ("1", "get-application-data", "ApplicationDataWorker", "GET /applications/{id}", "processType=SIGNING", "processInstanceKey, getProcessParams"),
            ("2", "transfer-control.v1", "внешний connector", "Kafka/connector", "messageName=SigningDocs, useProcessInstanceKey=true", "вне репы"),
            ("3", "message SigningDocs", "Zeebe receive", "correlationKey=processInstanceKey", "payload TBD", "продолжение процесса"),
        ],
    }
)

# PAYMENT
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-payment-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-payment-ncins-pa",
        "purpose": "BPMN payment + InsurancePaymentWorker (STUB hold/create) + ApplicationDataWorker. Timer в develop: PT15M.",
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "1) get-application-data (PAYMENT) → ApplicationDataWorker\n"
                    "2) ump-payment-ncins-pa.set-hold → InsurancePaymentWorker → PaymentService STUB\n"
                    "3) ump-payment-ncins-pa.create-payment → STUB\n"
                    "4) receive paymentFinished (corr=businessKey в develop)\n"
                    "5) Event SubProcess timer PT15M → timeoutMessage=TIMEOUT"
                ),
                "ump-payment-ncins-pa.bpmn + InsurancePaymentWorker",
                WHERE_ALGO,
                "Разработчик",
            ),
            ("2", "SLA по времени отклика", "Hold/create stub — мгновенно. Message почти мгновенно. Timer BPMN develop = PT15M.", "PT15M в develop", WHERE_SLA, "Аналитик"),
            ("3", "Частота планируемой нагрузки (ЧПН)", "≈85/час / пик ≈170 (после signing-timeout).", "15000×~95%", WHERE_CHPN, "Аналитик"),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                "[GET applications] ~85–170\n[set-hold STUB] ~85–170\n[create-payment STUB] ~85–170\n[paymentFinished message] ~85–170",
                "EQ HTTP в коде нет",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "get PAYMENT outs: fullName, productId, inn, contractNumber, insurancePremium, debitAccount.\n"
                    "set-hold / create-payment vars (IT): "
                    '{"businessKey":"550e8400-e29b-41d4-a716-446655110003","productCode":"NON_INSURANCE_CODE","systemCode":"SFA"}\n'
                    "HTTP EQ отсутствует — PaymentService TODO заглушка."
                ),
                "InsurancePaymentWorkerIT",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                (
                    "set-hold / create-payment: {\"code\":\"SUCCESS\"} (message null).\n"
                    "На ERROR маппер отдаёт code=ERROR + message.\n"
                    "Timeout: timeoutMessage=TIMEOUT (PT15M)."
                ),
                '{"code":"SUCCESS"}',
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "Stub ничего не пишет. UMP DB не меняется этим воркером.",
                "нет изменений",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "Hold/create уже stub в коде. paymentFinished — подложить Zeebe/Kafka message (corr=businessKey).",
                "PaymentService stub",
                WHERE_STUB,
                "Разработчик",
            ),
            ("11", "Ожидаемое время отклика замоканных внешних зависимостей", "Stub ~0; message 0–30с запас.", "мгновенно", WHERE_LAT, "Разработчик"),
            (
                "12",
                "Сценарий тестирования",
                "1) Success: get→set-hold→create-payment→paymentFinished.\n2) Timeout PT15M без message.\nПосле реальной EQ — обновить контракты.",
                "IT: InsurancePaymentWorkerIT",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "methods_section": "Методы payment (develop)",
        "method_rows": [
            ("1", "get-application-data", "ApplicationDataWorker", "GET /applications/{id}", "processType=PAYMENT", "fullName, productId, inn, contractNumber, insurancePremium, debitAccount"),
            ("2", "ump-payment-ncins-pa.set-hold", "InsurancePaymentWorker.setHold", "STUB PaymentService", "Map variables", '{"code":"SUCCESS"}'),
            ("3", "ump-payment-ncins-pa.create-payment", "InsurancePaymentWorker.createPayment", "STUB PaymentService", "Map variables", '{"code":"SUCCESS"}'),
            ("4", "message paymentFinished", "Zeebe receive", "corr=businessKey", "payload TBD", "End / continue"),
            ("5", "timer", "event sub-process", "PT15M", "—", "timeoutMessage=TIMEOUT"),
        ],
    }
)

# FINALISATION
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-finalisation-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-finalisation-ncins-pa",
        "purpose": "FinalisationWorker → POST /v1/ins-contracts + 3× ea-send-documents.v1 (внешний connector).",
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "1) get-application-data (FINALISATION)\n"
                    "2) ump-finalisation-ncins-pa.create-contract → FinalisationWorker → NibNcinsAccClient POST /v1/ins-contracts\n"
                    "3) Parallel Gateway → 3× ump-document-connectors.ea-send-documents.v1 "
                    "(acId=agreementLink|policyLink|contractLink, stopInIncident=false)\n"
                    "4) Join → End"
                ),
                "FinalisationWorker + NibNcinsAccClient",
                WHERE_ALGO,
                "Разработчик",
            ),
            ("2", "SLA по времени отклика", "Sync get + create-contract (+retries). 3 EA параллельно. Несколько секунд/вызов (TBD).", "несколько сек", WHERE_SLA, "Аналитик"),
            ("3", "Частота планируемой нагрузки (ЧПН)", "≈85/170 заявок; EA ×3 → ~255–510 вызовов connector/час в пике.", "×3 EA", WHERE_CHPN, "Аналитик"),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                "[GET applications] ~85–170\n[POST /v1/ins-contracts] ~85–170\n[ea-send-documents.v1 ×3] ~255–510",
                "integration.nib.corp-ncins-acc-api.url",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "HTTP: POST {nib.url}/v1/ins-contracts\n"
                    "Headers (IT): A-userId=AAAX22, A-customerId=AAAX22, A-clientType=BACKEND, A-channelId=ump "
                    "(+ optional A-userIp, A-projectId=ncins)\n"
                    "Camunda/ContractCreateInputVariables (из FinalisationWorkerIT.createValidVariables):\n"
                    "programId=3, inn=7707083893, email=romashka@rambler.ru, "
                    "contractNumber=Z6922/888/ABR14880/10, dates, sums, debitAccount, duration=12, "
                    "paymentType=payment_account, contract/policy/agreementLink, insuranceObjects[], "
                    "ownerId=AAAX22, phoneNumber, legalAddress, sellerId, sellerChannel=SFA\n"
                    "URL: http://corp-gateway-test.../corp-ncins-acc-api (yaml)"
                ),
                "FinalisationWorkerIT",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                (
                    "HTTP 201 {} (void).\n"
                    "Worker out success: {\"result\":\"SUCCESS\",\"errorMessage\":null}\n"
                    "5xx: 3 retry → {\"result\":\"ERROR\",\"errorMessage\":\"...\"}\n"
                    "401: refresh token + retry (IT).\n"
                    "EA response необязателен; stopInIncident=false."
                ),
                "201 + result=SUCCESS",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "UMP DB не меняется. Договор создаётся во внешней АБ (NIB). Документы — ЭА через connector.",
                "внешняя АБ + ЭА",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "NIB — стенд/WireMock (как IT). 3× ea-send — мок connector на НТ. Нужны уникальные contractNumber.",
                "WireMock POST /v1/ins-contracts",
                WHERE_STUB,
                "Разработчик",
            ),
            ("11", "Ожидаемое время отклика замоканных внешних зависимостей", "Несколько секунд; join ждёт 3 EA.", "несколько сек", WHERE_LAT, "Разработчик"),
            (
                "12",
                "Сценарий тестирования",
                "1) Success get→create→3 EA.\n2) create 503 → ERROR after 3 retries.\n3) 401 token refresh.\n4) EA без docs — не обязан инцидент.",
                "FinalisationWorkerIT",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "methods_section": "Методы finalisation (develop + IT)",
        "method_rows": [
            ("1", "get-application-data", "ApplicationDataWorker", "GET /applications/{id}", "processType=FINALISATION", "поля договора + agreement/policy/contractLink"),
            ("2", "ump-finalisation-ncins-pa.create-contract", "FinalisationWorker", "POST /v1/ins-contracts", "ContractRequestDto + headers A-*", "201 {}; worker result=SUCCESS|ERROR"),
            ("3", "ea-send ×3", "внешний connector", "ump-document-connectors.ea-send-documents.v1", "acId=agreement|policy|contractLink; stopInIncident=false", "вне репы"),
        ],
    }
)

# DELETE
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-delete-documents-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-delete-documents-ncins-pa",
        "purpose": "DeleteDocumentsWorker STUB + get-application-data DELETE_DOCS. BPMN multiInstance documents.",
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "1) get-application-data (DELETE_DOCS) → documents[{documentLink}] из contract/policy/agreement links\n"
                    "2) ump-delete-documents-ncins-pa.delete-documents multiInstance sequential "
                    "(inputCollection=documents, inputElement=document, outputCollection=acDocuments)\n"
                    "DeleteDocumentsService: TODO AC 1.0 — сейчас всегда status=1"
                ),
                "DeleteDocumentsWorker + BPMN",
                WHERE_ALGO,
                "Разработчик",
            ),
            ("2", "SLA по времени отклика", "Sync stub; время × N documents.", "несколько сек × N", WHERE_SLA, "Аналитик"),
            ("3", "Частота планируемой нагрузки (ЧПН)", "~5% заявок ≈4–5/час.", "15000×5%", WHERE_CHPN, "Аналитик"),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                "[GET applications] ~4–5\n[delete-documents STUB / будущий AC deleteDocsh] ~4–5 × N",
                "AC HTTP в коде ещё нет",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "Call Activity из main timeout (main BPMN в develop нет).\n"
                    "Worker input (IT):\n"
                    '{"businessKey":"550e8400-e29b-41d4-a716-446655110000",'
                    '"document":{"documentLink":"https://example.com/doc/123"}}'
                ),
                "DeleteDocumentsWorkerIT",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                (
                    "Success: {\"documentResponse\":{\"status\":\"1\"}}\n"
                    "Exception path: status=\"2\".\n"
                    "Реальный AC deleteDocsh — TBD (TODO в DeleteDocumentsService)."
                ),
                '{"documentResponse":{"status":"1"}}',
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "Сейчас stub — UMP DB не меняется. После AC — удаление во внешней AC.",
                "нет изменений UMP",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "Уже stub в коде. На НТ достаточно текущего мока до подключения AC.",
                "DeleteDocumentsService stub",
                WHERE_STUB,
                "Разработчик",
            ),
            ("11", "Ожидаемое время отклика замоканных внешних зависимостей", "Stub ~мгновенно.", "~0", WHERE_LAT, "Разработчик"),
            (
                "12",
                "Сценарий тестирования",
                "1) delete success status=1.\n2) multiInstance несколько documentLink.\n3) после AC — негатив невалидный UUID.",
                "DeleteDocumentsWorkerIT",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "methods_section": "Методы delete (develop + IT)",
        "method_rows": [
            ("1", "get-application-data", "ApplicationDataWorker", "GET /applications/{id}", "processType=DELETE_DOCS", "documents[{documentLink}]"),
            ("2", "ump-delete-documents-ncins-pa.delete-documents", "DeleteDocumentsWorker", "STUB (TODO AC 1.0)", "document.documentLink", '{"documentResponse":{"status":"1"}}'),
        ],
    }
)


def write_index() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Оглавление"
    ws["A1"] = "НТ пререквизиты ncins — по коду develop ump-ncins-pa@4a921a1653e"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Примеры request/response из *Worker.java + *WorkerIT + stubs. BPMN main-ma в develop нет."
    ws["A2"].font = SUB_FONT
    ws["A3"] = "Файл"
    ws["B3"] = "Процесс"
    style_header_row(ws, 3, 2)
    rows = [
        ("НТ пререквизиты __ Worker ump-main-ma-ncins-pa.xlsx", "Оркестратор (BPMN отсутствует в develop)"),
        ("НТ пререквизиты __ Worker ump-prepare-documents-ncins-pa.xlsx", "Подготовка документов"),
        ("НТ пререквизиты __ Worker ump-signing-documents-ncins-pa.xlsx", "Подписание"),
        ("НТ пререквизиты __ Worker ump-payment-ncins-pa.xlsx", "Оплата (stub + PT15M)"),
        ("НТ пререквизиты __ Worker ump-finalisation-ncins-pa.xlsx", "Финализация"),
        ("НТ пререквизиты __ Worker ump-delete-documents-ncins-pa.xlsx", "Удаление (stub)"),
    ]
    for i, (a, b) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=a).font = CELL_FONT
        ws.cell(row=i, column=2, value=b).font = CELL_FONT
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 45
    wb.save(OUT / "00-README-index.xlsx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_index()
    for meta in PROCESSES:
        build(meta).save(OUT / meta["file"])
        print("wrote", meta["file"])


if __name__ == "__main__":
    main()
