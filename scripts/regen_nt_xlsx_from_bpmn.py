#!/usr/bin/env python3
"""Regenerate all 6 NT prerequisite Excel workbooks from BPMN PR truth + SA/backend answers."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path("/workspace/nt-prerequisites-xlsx")
DATE = "UMP · 21 июля 2026 · актуализировано по BPMN PR (NCINS / commits 1–7)"

HEADER_FILL = PatternFill("solid", fgColor="D6EAF8")
TITLE_FONT = Font(name="Calibri", size=14, bold=True)
SUB_FONT = Font(name="Calibri", size=11, italic=True)
SECTION_FONT = Font(name="Calibri", size=12, bold=True)
HEADER_FONT = Font(name="Calibri", size=11, bold=True)
CELL_FONT = Font(name="Calibri", size=11)
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
COLS = ["#", "Пререквизит", "Значение", "Пример", "Где взять", "Предпочтительная роль заполняющего"]
COL_WIDTHS = {"A": 5, "B": 32, "C": 55, "D": 45, "E": 28, "F": 18}


def style_header_row(ws, row: int, ncols: int = 6) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = THIN


def write_title(ws, title: str, purpose: str | None = None) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = DATE
    ws["A2"].font = SUB_FONT
    if purpose:
        ws["A9" if ws.title == "Содержание" else "A3"] = purpose


def set_widths(ws) -> None:
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width


def write_toc(ws, title: str, purpose: str) -> None:
    write_title(ws, title)
    ws["A4"] = "Table of Contents"
    ws["A4"].font = SECTION_FONT
    ws["A5"] = "1 Согласование"
    ws["A6"] = "2 1. Минимальная информация для определения необходимости проведения НТ"
    ws["A7"] = (
        "3 2. Полная информация для проведения НТ "
        "(заполняется, только если получено заключение от инженера о необходимости проведения НТ)"
    )
    ws["A9"] = purpose
    ws["A9"].font = CELL_FONT
    ws["A9"].alignment = WRAP
    set_widths(ws)
    ws.row_dimensions[9].height = 45


def write_agreement(ws, title: str) -> None:
    write_title(ws, title)
    ws["A4"] = "1 Согласование"
    ws["A4"].font = SECTION_FONT
    ws["A5"] = "Аналитик"
    ws["B5"] = "Согласование"
    style_header_row(ws, 5, 2)
    ws["A6"] = "TBD"
    ws["B6"] = "TBD"
    for col in range(1, 3):
        ws.cell(row=6, column=col).font = CELL_FONT
        ws.cell(row=6, column=col).border = THIN
        ws.cell(row=6, column=col).alignment = WRAP
    set_widths(ws)


def write_table(ws, title: str, section: str, rows: list[tuple], start_header_row: int = 5) -> None:
    write_title(ws, title)
    ws.cell(row=4, column=1, value=section).font = SECTION_FONT
    for i, h in enumerate(COLS, 1):
        ws.cell(row=start_header_row, column=i, value=h)
    style_header_row(ws, start_header_row)
    for r_i, row in enumerate(rows, start_header_row + 1):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
        # taller rows for long values
        text = str(row[2]) if len(row) > 2 else ""
        ws.row_dimensions[r_i].height = min(220, max(60, 14 + text.count("\n") * 12))
    set_widths(ws)


def write_bpmn_sheet(ws, title: str, section: str, rows: list[tuple]) -> None:
    write_title(ws, title)
    ws.cell(row=4, column=1, value=section).font = SECTION_FONT
    headers = ["#", "BPMN id", "Name", "Type / called process", "Ключевые параметры", "STATIC / DYNAMIC"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=5, column=i, value=h)
    style_header_row(ws, 5)
    for r_i, row in enumerate(rows, 6):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
        ws.row_dimensions[r_i].height = 50
    set_widths(ws)
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 40


def build_workbook(meta: dict) -> Workbook:
    wb = Workbook()
    # TOC
    ws0 = wb.active
    ws0.title = "Содержание"
    write_toc(ws0, meta["title"], meta["purpose"])

    ws1 = wb.create_sheet("1. Согласование")
    write_agreement(ws1, meta["title"])

    ws2 = wb.create_sheet("2. Минимальная информация")
    write_table(
        ws2,
        meta["title"],
        "1. Минимальная информация для определения необходимости проведения НТ",
        meta["min_rows"],
    )

    ws3 = wb.create_sheet("3. Полная информация")
    write_table(
        ws3,
        meta["title"],
        "2. Полная информация для проведения НТ (заполняется, только если получено заключение от инженера о необходимости проведения НТ)",
        meta["full_rows"],
    )

    ws4 = wb.create_sheet("BPMN методы (источник)")
    write_bpmn_sheet(ws4, meta["title"], meta["bpmn_section"], meta["bpmn_rows"])
    return wb


# ---------------------------------------------------------------------------
# Process data
# ---------------------------------------------------------------------------

WHERE_ALGO = "Ссылка на алгоритм работы метода в git или confluence."
WHERE_SLA = "Получить от заказчика доработки. Если заказчик неизвестен, определяем экспертно."
WHERE_CHPN = "Аналитику нужно определить: 1) макс вызовов/час; 2) среднее вызовов/час; 3) способ получения оценки."
WHERE_CHPN6 = "Само НТ проводится по профилю прогнозируемой через полгода ЧПН."
WHERE_DEP = "Может совпадать с ЧПН."
WHERE_DEP6 = "Может совпадать с прогнозируемой через полгода ЧПН."
WHERE_CALL = "Сначала нужно добыть живой пример вызова сервиса (локально/стенд) и параметризовать."
WHERE_RESP = "Нужны примеры ответов, которые можно считать ожидаемыми."
WHERE_MUT = "Перечислить поля/таблицы/ресурсы, которые создаются/перезаписываются (кроме логирования)."
WHERE_MIG = "Опционально. Скрипт готовит разработчик только однажды перед боем."
WHERE_STUB = "Для зависимостей, которые нужно/можно глушить — прикрепить моки."
WHERE_LAT = "Потребуется для имитации нагрузки на внешние зависимости."
WHERE_SCEN = "Опционально. Готовит аналитик, если тестируется бизнес-процесс, а не метод."

PROCESSES: list[dict] = []

# ---- MAIN ----
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-main-ma-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-main-ma-ncins-pa",
        "purpose": (
            "Назначение: Управление процессом мультизаявки по некредитному страхованию. "
            "Источник правды: BPMN ump-main-ma-ncins-pa.bpmn (PR NCINS)."
        ),
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "[ump-main-ma-ncins-pa] Оркестратор мультизаявки ncins (BPMN).\n"
                    "Цепочка:\n"
                    "1) Start (businessKey DYNAMIC, productCode≈NON_CREDIT_INSURANCE)\n"
                    "2) Call Activity ump-prepare-documents-ncins-pa\n"
                    "3) Kafka stage DOCS_COMPLETED (kafka-product-stage-outbound-connector)\n"
                    "4) Call Activity ump-signing-documents-ncins-pa\n"
                    "5) XOR: если timeoutMessage=\"TIMEOUT\" → Call Activity ump-delete-documents-ncins-pa "
                    "→ product-status BANK_REJECT; иначе stage SIGN_COMPLETED\n"
                    "6) Call Activity ump-payment-ncins-pa\n"
                    "7) XOR: TIMEOUT → delete + BANK_REJECT; иначе stage PAYMENT_COMPLETED\n"
                    "8) Call Activity ump-finalisation-ncins-pa → stage COMPLETED → End result=SUCCESS\n"
                    "Event SubProcess: Timer PT30M → Call Activity ump-delete-documents-ncins-pa → BANK_REJECT "
                    "(forcefullyTerminate=true)."
                ),
                "Файл: bpmn/.../ump-main-ma-ncins-pa.bpmn",
                WHERE_ALGO,
                "Разработчик",
            ),
            (
                "2",
                "SLA по времени отклика",
                (
                    "E2E SLA main — TBD на встрече с НТ. Sync-шаги дочерних: несколько секунд/вызов (ориентир СА). "
                    "Таймеры BPMN: signing PT25M, payment PT5M, main emergency PT30M."
                ),
                "TBD числом с инженером НТ",
                WHERE_SLA,
                "Аналитик",
            ),
            (
                "3",
                "Частота планируемой нагрузки (ЧПН)",
                (
                    "База 15000 заявок/мес (SFA, рабочие часы ЮЛ). Среднее ≈89/час; пик (×2) ≈180/час. "
                    "Профиль: ~95% happy-path / ~5% timeout (из timeout ~95% — signing)."
                ),
                "15000/~21 р.д./8 ч ≈ 89/час; пик ×2",
                WHERE_CHPN,
                "Аналитик",
            ),
            (
                "4",
                "Прогнозируемая через полгода ЧПН",
                "TBD (прогноз через 6 мес не зафиксирован командой).",
                "TBD макс/среднее через 6 мес",
                WHERE_CHPN6,
                "Аналитик",
            ),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                (
                    "Сервис/Процесс | ЧПН ср./пик (вызовов/час):\n"
                    "[ump-prepare-documents-ncins-pa] ~89 / ~180\n"
                    "[ump-signing-documents-ncins-pa] ~89 / ~180\n"
                    "[ump-payment-ncins-pa] ~85 / ~170\n"
                    "[ump-finalisation-ncins-pa] ~85 / ~170\n"
                    "[ump-delete-documents-ncins-pa] ~4–5 / ~9 (только timeout)\n"
                    "[kafka-product-stage-outbound-connector] ~89–180 (стадии DOCS/SIGN/PAYMENT/COMPLETED)\n"
                    "[kafka-product-status-event-outbound-connector] ~4–5 (BANK_REJECT)\n"
                    "Kafka signing/payment flow: ориентир до ~1800 msg/день (бэкенд)."
                ),
                "Happy-path = ЧПН оркестратора; delete только timeout",
                WHERE_DEP,
                "Аналитик",
            ),
            (
                "6",
                "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                "TBD (пропорционально п.4). Пока ориентир = п.5.",
                "TBD",
                WHERE_DEP6,
                "Аналитик",
            ),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "HTTP-эндпоинта нет. Старт через Camunda/Zeebe.\n"
                    "Вход: businessKey (DYNAMIC UUID), productCode=NON_CREDIT_INSURANCE (STATIC для ncins).\n"
                    "Map.of(\"businessKey\", \"<uuid>\", \"productCode\", \"NON_CREDIT_INSURANCE\");\n"
                    "client.newCreateInstanceCommand().bpmnProcessId(\"ump-main-ma-ncins-pa\")..."
                ),
                "Параметризация: набор businessKey ≥ ЧПН; productCode — const",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                (
                    "Happy-path End: result=SUCCESS.\n"
                    "Reject ends: statusCode=BANK_REJECT, statusMessage=\"Превышено время ожидания\", "
                    "forcefullyTerminate=true, cancelUserTask=false.\n"
                    "SUCCESS rate ориентир 99.99% — TBD с НТ."
                ),
                "SUCCESS / BANK_REJECT",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                (
                    "Main напрямую UMP-БД не меняет. Стадии/статусы через Kafka-коннекторы: "
                    "DOCS_COMPLETED → SIGN_COMPLETED → PAYMENT_COMPLETED → COMPLETED / BANK_REJECT."
                ),
                "Стадии продукта через outbound connectors",
                WHERE_MUT,
                "Разработчик",
            ),
            (
                "9",
                "Скрипт, содержащий сопутствующие изменения в БД (миграция)",
                "Нет необходимости",
                "—",
                WHERE_MIG,
                "Разработчик",
            ),
            (
                "10",
                "Заглушки для внешних зависимостей",
                (
                    "Моки в дочерних: EQ (payment), ЭА (finalisation, 3× ea-send), AC delete (пока мок).\n"
                    "На НТ обязательно подкладывать Kafka completion для signing (SigningDocs) и payment "
                    "(paymentFinished) с correlationKey = businessKey+\".\"+\"NON_CREDIT_INSURANCE\"."
                ),
                "Моки дочерних + Kafka publish completion",
                WHERE_STUB,
                "Разработчик",
            ),
            (
                "11",
                "Ожидаемое время отклика замоканных внешних зависимостей",
                (
                    "Sync дочерних: несколько секунд/вызов. Async signing на НТ: несколько секунд, ≤~3 (TBD). "
                    "Payment message: почти мгновенно (+запас ~30с). Kafka: запас по ~1800 msg/день."
                ),
                "sync ~несколько сек; kafka по стенду",
                WHERE_LAT,
                "Разработчик",
            ),
            (
                "12",
                "Сценарий тестирования",
                (
                    "1) Happy-path E2E (~95%): prepare→DOCS_COMPLETED→signing→SIGN_COMPLETED→payment→"
                    "PAYMENT_COMPLETED→finalisation→COMPLETED→SUCCESS.\n"
                    "2) Signing timeout (~4.75%): timeoutMessage=TIMEOUT→Call Activity delete→BANK_REJECT.\n"
                    "3) Payment timeout (малая): то же через XOR после payment.\n"
                    "4) Main emergency PT30M (~0): delete→BANK_REJECT.\n"
                    "Корреляция async: businessKey.\".\"NON_CREDIT_INSURANCE."
                ),
                "Последовательные Call Activity + доля timeout 5%",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "bpmn_section": "Шаги из ump-main-ma-ncins-pa.bpmn — STATIC vs DYNAMIC",
        "bpmn_rows": [
            ("1", "StartEvent", "Start", "startEvent", "businessKey, productCode", "DYNAMIC: businessKey; STATIC≈ productCode=NON_CREDIT_INSURANCE"),
            ("2", "call-activity-prepare", "Prepare docs", "calledElement: ump-prepare-documents-ncins-pa", "businessKey, productCode", "DYNAMIC in; out: result"),
            ("3", "connector-stage-docs", "Stage DOCS_COMPLETED", "kafka-product-stage-outbound-connector", "stageCode=DOCS_COMPLETED", "STATIC: stage; DYNAMIC: businessKey/productCode"),
            ("4", "call-activity-signing", "Signing", "calledElement: ump-signing-documents-ncins-pa", "businessKey, productCode", "DYNAMIC; out: result / timeoutMessage"),
            ("5", "Gateway_13ehqon", "XOR after signing", "exclusiveGateway", "timeoutMessage=\"TIMEOUT\" → delete", "DYNAMIC: timeoutMessage"),
            ("6", "call-activity-delete*", "Delete docs", "calledElement: ump-delete-documents-ncins-pa", "на timeout-ветках (+ PT30M)", "DYNAMIC: businessKey, productCode"),
            ("7", "connector-stage-sign", "Stage SIGN_COMPLETED", "kafka-product-stage-outbound-connector", "SIGN_COMPLETED", "STATIC stage"),
            ("8", "call-activity-payment", "Payment", "calledElement: ump-payment-ncins-pa", "businessKey, productCode", "DYNAMIC; out: result / timeoutMessage"),
            ("9", "Gateway_1j90rp8", "XOR after payment", "exclusiveGateway", "TIMEOUT → delete", "DYNAMIC: timeoutMessage"),
            ("10", "connector-stage-pay", "Stage PAYMENT_COMPLETED", "kafka-product-stage-outbound-connector", "PAYMENT_COMPLETED", "STATIC stage"),
            ("11", "call-activity-finalisation", "Finalisation", "calledElement: ump-finalisation-ncins-pa", "businessKey, productCode", "DYNAMIC"),
            ("12", "connector-stage-done", "Stage COMPLETED", "kafka-product-stage-outbound-connector", "COMPLETED", "STATIC stage"),
            ("13", "EventSubProcess timer", "Emergency timeout", "timerEventDefinition PT30M", "→ delete → BANK_REJECT", "STATIC: PT30M"),
            ("14", "status-reject", "BANK_REJECT", "kafka-product-status-event-outbound-connector", "statusCode=BANK_REJECT; forcefullyTerminate=true", "STATIC status fields"),
        ],
    }
)

# ---- PREPARE ----
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-prepare-documents-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-prepare-documents-ncins-pa",
        "purpose": (
            "Назначение: Процесс подготовки документов (ncins). "
            "Источник правды: BPMN ump-prepare-documents-ncins-pa.bpmn (PR NCINS)."
        ),
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "[ump-prepare-documents-ncins-pa] Процесс подготовки документов (BPMN).\n"
                    "Цепочка:\n"
                    "1) StartEvent\n"
                    "2) Service Task get-application-data — processType=PREPARE_DOCUMENTS (STATIC), businessKey (DYNAMIC)\n"
                    "3) Service Task ump-prepare-documents-ncins-pa.get-report-data → documents/reportData\n"
                    "4) Call Activity ump-generate-and-save-document-pa "
                    "(serviceId=businessKey, serviceCode=\"APPLICATION\" STATIC, productCode, documents)\n"
                    "5) Service Task update-product → End\n"
                    "MultiInstance на generate в финальном BPMN нет (Call Activity)."
                ),
                "Файл: bpmn/.../ump-prepare-documents-ncins-pa.bpmn",
                WHERE_ALGO,
                "Разработчик",
            ),
            (
                "2",
                "SLA по времени отклика",
                "Все шаги синхронные (service task / call activity). Ориентир СА: несколько секунд на вызов. Точное число — TBD с НТ.",
                "несколько секунд / sync-вызов (TBD мс)",
                WHERE_SLA,
                "Аналитик",
            ),
            (
                "3",
                "Частота планируемой нагрузки (ЧПН)",
                (
                    "От общей базы 15000 заявок/мес (SFA). Среднее ≈89/час; пик ≈180/час. "
                    "1 документ на заявку. Пример размера reportData ≈283 КБ (avg/p95/max — TBD)."
                ),
                "15000/~21/8 ≈ 89/час; пик ×2",
                WHERE_CHPN,
                "Аналитик",
            ),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                (
                    "[get-application-data] ~89 / ~180\n"
                    "[ump-prepare-documents-ncins-pa.get-report-data] ~89 / ~180\n"
                    "[ump-generate-and-save-document-pa] (ПФ+AC) ~89 / ~180\n"
                    "[update-product] ~89 / ~180"
                ),
                "Каждый шаг 1 раз на заявку",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD (ориентир = п.5).", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "HTTP нет. Запуск Call Activity из main или напрямую в Zeebe.\n"
                    "Старт vars: businessKey (DYNAMIC), productCode.\n"
                    "Внутри: processType=PREPARE_DOCUMENTS (STATIC); Call Activity: serviceCode=\"APPLICATION\" (STATIC FEEL)."
                ),
                "Параметризация businessKey ≥ ЧПН",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                (
                    "Успех: все task SUCCESS → End.\n"
                    "Call Activity outs: docsResult, acRequestId, acDocuments.\n"
                    "Ошибки ПФ/AC: HTTP 400/500 (полный каталог — TBD у Жени)."
                ),
                "SUCCESS / ERROR; 400; 500",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "update-product: фактически значимо contractLink (PUT продукта). Заявка читается в get-application-data.",
                "contractLink (+ поля product из маппинга)",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "Внешний вызов через Call Activity ump-generate-and-save-document-pa (ПФ + AlfaCapture). На НТ без своей папки AC реальное сохранение ограничено — уточнять мок/стенд.",
                "Мок/стенд generate-and-save при необходимости",
                WHERE_STUB,
                "Разработчик",
            ),
            (
                "11",
                "Ожидаемое время отклика замоканных внешних зависимостей",
                "Как у sync: несколько секунд на вызов (TBD), включая call activity.",
                "get-report-data / generate-and-save / update-product — по несколько сек",
                WHERE_LAT,
                "Разработчик",
            ),
            (
                "12",
                "Сценарий тестирования",
                (
                    "1) Start → get-application-data (PREPARE_DOCUMENTS) → get-report-data → "
                    "Call Activity generate-and-save → update-product → End.\n"
                    "2) Ошибка get 4xx/5xx.\n"
                    "3) Ошибка генерации ПФ 400/500.\n"
                    "4) Пик ~180/час, лимит потоков 200."
                ),
                "Последовательный вызов 4 исполняемых шагов",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "bpmn_section": "Шаги из ump-prepare-documents-ncins-pa.bpmn — STATIC vs DYNAMIC",
        "bpmn_rows": [
            ("1", "StartEvent_1", "Start", "startEvent", "ожидает vars процесса", "DYNAMIC_START: businessKey, productCode"),
            ("2", "service-task-get-params", "Получить данные по заявке", "zeebe type: get-application-data", "processType → PREPARE_DOCUMENTS", "STATIC: processType=PREPARE_DOCUMENTS; DYNAMIC: businessKey"),
            ("3", "service-task-get-report-data", "Подготовить данные для ПФ", "type: ump-prepare-documents-ncins-pa.get-report-data", "формирует documents/reportData", "DYNAMIC_FROM_APP → documents"),
            ("4", "call-activity-ump-generate-and-save-document-pa", "Формирование и сохранение в AlfaCapture", "calledElement: ump-generate-and-save-document-pa", "serviceId=businessKey; serviceCode=\"APPLICATION\"; productCode; documents; out: docsResult, acRequestId, acDocuments", "STATIC: serviceCode; DYNAMIC: serviceId, productCode, documents"),
            ("5", "service-task-update-product", "Обновить данные по продукту", "type: update-product", "после AC", "DYNAMIC_FROM_PREV → product update"),
            ("6", "End", "End", "endEvent", "result=SUCCESS", "—"),
        ],
    }
)

# ---- SIGNING ----
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-signing-documents-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-signing-documents-ncins-pa",
        "purpose": (
            "Назначение: Подписание документов. "
            "Источник правды: BPMN ump-signing-documents-ncins-pa.bpmn (PR NCINS). Timer PT25M."
        ),
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "[ump-signing-documents-ncins-pa] (BPMN):\n"
                    "1) Start → get-application-data (processType=SIGNING STATIC)\n"
                    "2) Connector ump-process-transfer-connector → Kafka ump.process.to.system "
                    "(systemCode=B2B STATIC, messageName=GenerateDocs STATIC, stopInIncident=true; "
                    "applicationId=businessKey; correlationKey=businessKey+\".\"+productCode)\n"
                    "3) kafka-product-stage-outbound-connector stage=SIGN_WAITING (STATIC)\n"
                    "4) Receive Message SigningDocs; correlationKey = businessKey+\".\"+\"NON_CREDIT_INSURANCE\" (STATIC product в FEEL)\n"
                    "5) update-product → End SUCCESS\n"
                    "Event SubProcess Timer PT25M → status BANK_REJECT (forcefullyTerminate=true).\n"
                    "Delete Call Activity в signing нет — cleanup делает main по timeoutMessage."
                ),
                "BPMN + Confluence ump.process.to.system / from.system",
                WHERE_ALGO,
                "Разработчик",
            ),
            (
                "2",
                "SLA по времени отклика",
                "Sync: несколько секунд. Async в проде до PT25M; для НТ сократить: несколько секунд, ≤~3 (TBD с НТ).",
                "НТ async wait ≤~3 сек; прод PT25M",
                WHERE_SLA,
                "Аналитик",
            ),
            (
                "3",
                "Частота планируемой нагрузки (ЧПН)",
                "≈89/час среднее, ≈180/час макс (от 15000/мес). Дублей сообщений по задумке быть не должно.",
                "15000/мес",
                WHERE_CHPN,
                "Аналитик",
            ),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                (
                    "[get-application-data] ~89–180/час\n"
                    "[ump-process-transfer-connector / ump.process.to.system] ~89–180/час\n"
                    "[kafka-product-stage SIGN_WAITING] ~89–180/час\n"
                    "[Receive SigningDocs / ump.process.from.system] ~89–180/час (на НТ подкладываем)\n"
                    "[update-product] ~89–180/час"
                ),
                "Ссылки Confluence to.system / from.system",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "Старт из main Call Activity: businessKey, productCode=NON_CREDIT_INSURANCE.\n"
                    "На НТ после to.system опубликовать SigningDocs в from.system с "
                    "correlationKey = \"<businessKey>.NON_CREDIT_INSURANCE\".\n"
                    "Live JSON примеров сообщений — TBD аналитика (подписание моковое)."
                ),
                "Kafka publish; correlationKey=businessKey.NON_CREDIT_INSURANCE",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                (
                    "Success: result=SUCCESS → main идёт в оплату (SIGN_COMPLETED).\n"
                    "Timeout PT25M: BANK_REJECT; main получает timeoutMessage и вызывает delete.\n"
                    "Late message после завершения — игнор; дубликатов быть не должно."
                ),
                "SUCCESS vs BANK_REJECT",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "Стадия SIGN_WAITING через Kafka; update-product (agreementLink из сообщения).",
                "stage + agreementLink",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "HTTP signing-канал не мокается. На НТ подкладываем completion в Kafka (вручную; автоматизация — с НТ). JSON-контракт — TBD.",
                "Инструкция publish SigningDocs",
                WHERE_STUB,
                "Разработчик",
            ),
            (
                "11",
                "Ожидаемое время отклика замоканных внешних зависимостей",
                "Не симулируем 1–25 мин. Completion на НТ: несколько секунд, ≤~3 (TBD).",
                "≤3 сек на НТ",
                WHERE_LAT,
                "Разработчик",
            ),
            (
                "12",
                "Сценарий тестирования",
                (
                    "1) Success: get→to.system→SIGN_WAITING→подложить SigningDocs "
                    "(corr=businessKey.NON_CREDIT_INSURANCE)→update-product→SUCCESS.\n"
                    "2) Timeout: не публиковать → PT25M → BANK_REJECT; main → delete.\n"
                    "3) Late message — игнор."
                ),
                "Normal + timeout",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "bpmn_section": "Шаги из ump-signing-documents-ncins-pa.bpmn — STATIC vs DYNAMIC",
        "bpmn_rows": [
            ("1", "Start", "Start", "startEvent", "businessKey, productCode", "DYNAMIC start"),
            ("2", "service-task-get-params", "Получить данные", "get-application-data", "processType=SIGNING", "STATIC processType; DYNAMIC businessKey"),
            ("3", "transfer-connector", "Передать управление", "ump-process-transfer-connector → ump.process.to.system", "systemCode=B2B; messageName=GenerateDocs; stopInIncident=true", "STATIC codes; DYNAMIC applicationId/correlationKey"),
            ("4", "stage-connector", "SIGN_WAITING", "kafka-product-stage-outbound-connector", "stage=SIGN_WAITING", "STATIC stage"),
            ("5", "message-catch", "SigningDocs", "message intermediateCatch", "correlationKey=businessKey+\".\"+\"NON_CREDIT_INSURANCE\"", "STATIC product in FEEL; DYNAMIC businessKey"),
            ("6", "update-product", "Обновить продукт", "type: update-product", "agreementLink из сообщения", "DYNAMIC"),
            ("7", "EventSubProcess", "Timeout", "timer PT25M → BANK_REJECT", "forcefullyTerminate=true", "STATIC PT25M"),
        ],
    }
)

# ---- PAYMENT ----
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-payment-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-payment-ncins-pa",
        "purpose": (
            "Назначение: Оплата страховки. "
            "Источник правды: BPMN ump-payment-ncins-pa.bpmn (PR NCINS). Timer PT5M. EQ — MOCK/DRAFT."
        ),
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "[ump-payment-ncins-pa] (BPMN):\n"
                    "1) Start → get-application-data (businessKey DYNAMIC; processType контекст PAYMENT)\n"
                    "2) Job ump-payment-ncins-pa.set-hold — DRAFT/MOCK\n"
                    "3) Job ump-payment-ncins-pa.create-payment — DRAFT/MOCK\n"
                    "4) Receive Message paymentFinished (Kafka ump.process.from.system); "
                    "correlationKey = businessKey+\".\"+\"NON_CREDIT_INSURANCE\"\n"
                    "5) End SUCCESS\n"
                    "Event SubProcess Timer PT5M → BANK_REJECT.\n"
                    "Delete при timeout выполняет main (XOR по timeoutMessage)."
                ),
                "BPMN; запуск из main после signing",
                WHERE_ALGO,
                "Разработчик",
            ),
            (
                "2",
                "SLA по времени отклика",
                "hold/create — sync несколько секунд (мок). message-wait — почти мгновенно; запас СА ~30 сек. Таймер BPMN PT5M.",
                "sync несколько сек; async ~0–30с; timer PT5M",
                WHERE_SLA,
                "Аналитик",
            ),
            (
                "3",
                "Частота планируемой нагрузки (ЧПН)",
                "≈85/час среднее, ≈170/час макс (15000×~95% после signing-timeout). 1 инстанс = 1 заявка.",
                "15000×~95%",
                WHERE_CHPN,
                "Аналитик",
            ),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                (
                    "[get-application-data] ~85–170/час\n"
                    "[EQ hold / create-payment MOCK] ~85–170/час\n"
                    "[Kafka paymentFinished] ~85–170/час (подкладываем на НТ)"
                ),
                "EQ + Kafka",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "Включается из main после signing. Вход: businessKey, productCode=NON_CREDIT_INSURANCE.\n"
                    "Данные заявки (get): fullName, productId, inn, contractNumber, insurancePremium, debitAccount.\n"
                    "Мок EQ: code=SUCCESS, message=null.\n"
                    "Подложить paymentFinished с correlationKey=\"<businessKey>.NON_CREDIT_INSURANCE\"."
                ),
                "get → 2 мока → Kafka paymentFinished",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                "Мок: code=SUCCESS, message=null. Негативы EQ 4xx/5xx пока не обязательны (mock-only). Timeout — PT5M → BANK_REJECT.",
                "SUCCESS / BANK_REJECT",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "Отдельной параметризации нет. EQ мок в UMP ничего не пишет. Дублей быть не может.",
                "Существенных изменений UMP DB не описывалось",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "EQ полностью замокан. Kafka completion подкладываем (в проде слал бы EQ). Живой контракт EQ сейчас неактуален.",
                "Мок EQ + Kafka paymentFinished",
                WHERE_STUB,
                "Разработчик",
            ),
            (
                "11",
                "Ожидаемое время отклика замоканных внешних зависимостей",
                "Hold/create mock: несколько секунд. Payment message: почти мгновенно (+запас 30с).",
                "мок EQ ~несколько сек; message 0–30с",
                WHERE_LAT,
                "Разработчик",
            ),
            (
                "12",
                "Сценарий тестирования",
                (
                    "1) Success: get→mock hold→mock create→подложить paymentFinished "
                    "(corr=businessKey.NON_CREDIT_INSURANCE)→SUCCESS.\n"
                    "2) Timeout: не публиковать → PT5M → BANK_REJECT; main → delete."
                ),
                "Основной + timeout",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "bpmn_section": "Шаги из ump-payment-ncins-pa.bpmn — STATIC vs DYNAMIC",
        "bpmn_rows": [
            ("1", "Start", "Start", "startEvent", "businessKey, productCode", "DYNAMIC start"),
            ("2", "service-task-get-params", "Получить данные", "get-application-data", "данные для оплаты", "DYNAMIC businessKey"),
            ("3", "set-hold", "Установить холд", "ump-payment-ncins-pa.set-hold", "DRAFT/MOCK EQ", "MOCK"),
            ("4", "create-payment", "Создать ПП", "ump-payment-ncins-pa.create-payment", "DRAFT/MOCK EQ", "MOCK"),
            ("5", "message-catch", "paymentFinished", "message intermediateCatch / ump.process.from.system", "correlationKey=businessKey+\".\"+\"NON_CREDIT_INSURANCE\"", "STATIC product in FEEL; DYNAMIC businessKey"),
            ("6", "EventSubProcess", "Timeout", "timer PT5M → BANK_REJECT", "forcefullyTerminate", "STATIC PT5M"),
        ],
    }
)

# ---- FINALISATION ----
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-finalisation-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-finalisation-ncins-pa",
        "purpose": (
            "Назначение: Финализация заявки. "
            "Источник правды: BPMN ump-finalisation-ncins-pa.bpmn (PR NCINS). "
            "После create-contract — Parallel Gateway и 3× ea-send-documents.v1."
        ),
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "[ump-finalisation-ncins-pa] (BPMN FINAL):\n"
                    "1) Start → get-application-data (processType=FINALISATION STATIC)\n"
                    "2) Service Task ump-finalisation-ncins-pa.create-contract (POST /v1/ins-contracts)\n"
                    "3) Parallel Gateway split → 3× Connector ump-document-connectors.ea-send-documents.v1:\n"
                    "   • acId=agreementLink, stopInIncident=false\n"
                    "   • acId=policyLink, stopInIncident=false\n"
                    "   • acId=contractLink, stopInIncident=false\n"
                    "4) Parallel Gateway join → End\n"
                    "MultiInstance на EA в финальном BPMN нет (было в ранних версиях)."
                ),
                "BPMN PR: 3 parallel ea-send",
                WHERE_ALGO,
                "Разработчик",
            ),
            (
                "2",
                "SLA по времени отклика",
                "Синхронные вызовы. Несколько секунд на каждый sync-вызов (TBD точное число с НТ). 3 EA идут параллельно.",
                "несколько секунд / вызов",
                WHERE_SLA,
                "Аналитик",
            ),
            (
                "3",
                "Частота планируемой нагрузки (ЧПН)",
                "≈85/час среднее, ≈170/час макс (дошедшие до финала от 15000/мес). На EA-коннектор ×3 вызова на заявку → ~255–510/час пик на connector type.",
                "15000×~95%; EA ×3",
                WHERE_CHPN,
                "Аналитик",
            ),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                (
                    "[get-application-data] ~85–170/час\n"
                    "[ump-finalisation-ncins-pa.create-contract / POST /v1/ins-contracts] ~85–170/час\n"
                    "[ump-document-connectors.ea-send-documents.v1] ×3 параллельно ≈255–510/час на тип "
                    "(agreementLink / policyLink / contractLink) — МОК на НТ"
                ),
                "АБ + 3× EA/AC",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "Старт: businessKey + processType=FINALISATION (STATIC).\n"
                    "Create contract request (пример): inn, email, contractNumber (НЕ статический — "
                    "из /v1/ins-contracts/contract-number), dates, sums, paymentType, agreementLink/contractLink, "
                    "insuranceObjects, ownerId, phone, legalAddress, sellerId, sellerChannel=SFA, programId.\n"
                    "EA: 3 вызова с acId из get-application; stopInIncident=false.\n"
                    "На int нагенерить список contractNumber для НТ; нужна интеграция стенда с контуром договоров."
                ),
                "Разные contractNumber на каждый прогон",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                (
                    "Get: getProcessParams=SUCCESS|ERROR.\n"
                    "Create contract: result=SUCCESS, errorMessage=null | result=ERROR + errorMessage.\n"
                    "EA Response: необязателен; при отсутствии документов не обязан инцидент (stopInIncident=false)."
                ),
                "SUCCESS/ERROR по таскам",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "БД UMP не изменяется. Договор — во внешней АБ; документы — в ЭА (на НТ мок).",
                "UMP DB без изменений",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "Мок на 3× ea-send-documents.v1. АБ ins-contracts — живой/стендовый. stopInIncident=false.",
                "Мок EA ×3; стенд POST /v1/ins-contracts",
                WHERE_STUB,
                "Разработчик",
            ),
            (
                "11",
                "Ожидаемое время отклика замоканных внешних зависимостей",
                "Sync: несколько секунд (TBD). Мок ЭА — в том же бюджете; параллельный join ждёт все 3.",
                "несколько сек",
                WHERE_LAT,
                "Разработчик",
            ),
            (
                "12",
                "Сценарий тестирования",
                (
                    "1) Success: get→create-contract→3× mock EA (parallel)→join→SUCCESS.\n"
                    "2) get ERROR.\n"
                    "3) create-contract ERROR.\n"
                    "4) EA без документов — не обязан инцидент.\n"
                    "Повторяемый прогон: уникальный contractNumber."
                ),
                "create-contract + 3 parallel ea-send",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "bpmn_section": "Шаги из ump-finalisation-ncins-pa.bpmn — STATIC vs DYNAMIC",
        "bpmn_rows": [
            ("1", "Start", "Start", "startEvent", "businessKey", "DYNAMIC"),
            ("2", "service-task-get-params", "Получить данные", "get-application-data", "processType=FINALISATION; outs: agreementLink, policyLink, contractLink, …", "STATIC processType; DYNAMIC businessKey"),
            ("3", "service-task-create-contract", "Создать договор", "ump-finalisation-ncins-pa.create-contract", "POST /v1/ins-contracts; out result/errorMessage", "DYNAMIC payload"),
            ("4", "Gateway_09hug5l", "Parallel split", "parallelGateway", "3 ветки EA", "—"),
            ("5", "send-documents-to-EA1", "EA agreement", "ump-document-connectors.ea-send-documents.v1", "acId=agreementLink; stopInIncident=false; retries=1", "STATIC stopInIncident; DYNAMIC acId"),
            ("6", "send-documents-to-EA", "EA policy", "ump-document-connectors.ea-send-documents.v1", "acId=policyLink; stopInIncident=false", "STATIC stopInIncident; DYNAMIC acId"),
            ("7", "send-documents-to-EA2", "EA contract", "ump-document-connectors.ea-send-documents.v1", "acId=contractLink; stopInIncident=false", "STATIC stopInIncident; DYNAMIC acId"),
            ("8", "Gateway_1sccgl0", "Parallel join", "parallelGateway", "ожидание 3 веток", "—"),
            ("9", "End", "End", "endEvent", "—", "—"),
        ],
    }
)

# ---- DELETE ----
PROCESSES.append(
    {
        "file": "НТ пререквизиты __ Worker ump-delete-documents-ncins-pa.xlsx",
        "title": "НТ пререквизиты :: Worker ump-delete-documents-ncins-pa",
        "purpose": (
            "Назначение: Удаление документов. BPMN process id: ump-delete-documents-ncins-pa "
            "(в Confluence/архиве часто ump-delete-documents-ncins-pa-ncins-pa). "
            "Вызов только из timeout-веток main. multiInstance по documents."
        ),
        "min_rows": [
            (
                "1",
                "Описание алгоритма работы тестируемого метода",
                (
                    "[ump-delete-documents-ncins-pa] (BPMN):\n"
                    "1) Start (businessKey DYNAMIC, productCode)\n"
                    "2) get-application-data — processType=DELETE_DOCS (STATIC FEEL =DELETE_DOCS)\n"
                    "3) Service Task ump-delete-documents-ncins-pa.delete-documents — multiInstance sequential:\n"
                    "   inputCollection=documents, inputElement=document,\n"
                    "   outputCollection=acDocuments, outputElement=documentResponse\n"
                    "   (прокси POST AlfaCapture deleteDocsh, syscode=UMP)\n"
                    "4) End\n"
                    "Источники вызова (Call Activity из main): signing timeout / payment timeout / main PT30M.\n"
                    "Сейчас таска в BPMN моковая."
                ),
                "BPMN id ump-delete-documents-ncins-pa + Confluence deleteDocsh",
                WHERE_ALGO,
                "Разработчик",
            ),
            (
                "2",
                "SLA по времени отклика",
                "Sync. Ориентир: несколько секунд (TBD). multiInstance sequential — время × число documents.",
                "несколько секунд × N docs",
                WHERE_SLA,
                "Аналитик",
            ),
            (
                "3",
                "Частота планируемой нагрузки (ЧПН)",
                "~5% от всех заявок (оценка сверху): ≈4–5/час среднее, ≈9/час макс. Из timeout ~95% — signing (PT25M).",
                "15000×5%",
                WHERE_CHPN,
                "Аналитик",
            ),
            ("4", "Прогнозируемая через полгода ЧПН", "TBD", "TBD", WHERE_CHPN6, "Аналитик"),
            (
                "5",
                "ЧПН на вызываемые внутри тестируемого метода другие методы сервисов",
                (
                    "[get-application-data] ~4–5/час\n"
                    "[AlfaCapture 1.0 POST deleteDocsh via delete-documents] ~4–5/час × N documents (multiInstance)\n"
                    "https://confluence.moscow.alfaintra.net/spaces/SMP/pages/2248960447/..."
                ),
                "Единственная внешняя зависимость — AC",
                WHERE_DEP,
                "Аналитик",
            ),
            ("6", "Прогнозируемая через полгода ЧПН на вызываемые внутри тестируемого метода другие методы сервисов", "TBD", "TBD", WHERE_DEP6, "Аналитик"),
        ],
        "full_rows": [
            (
                "6",
                "Пример кода вызова тестируемого метода сервиса",
                (
                    "Call Activity из main timeout-веток. Вход: businessKey, productCode=NON_CREDIT_INSURANCE.\n"
                    "documents[] из заявки (contractLink/policyLink/agreementLink). syscode=UMP (STATIC).\n"
                    "На НТ без папки AC — мок. Контракт deleteDocsh — TBD аналитика."
                ),
                "Call Activity; мок delete",
                WHERE_CALL,
                "Тестировщик",
            ),
            (
                "7",
                "Примеры кода ответа тестируемого метода сервиса",
                "В Camunda: acDocuments (outputCollection), элемент вида {\"status\":\"1\"}. Ошибка: не UUID документа. HTTP codes — TBD.",
                "[{\"status\":\"1\"}]",
                WHERE_RESP,
                "Разработчик",
            ),
            (
                "8",
                "Изменяемые данные в рамках работы метода",
                "В БД UMP изменений нет: только AC. Аудит UMP не затрагивается.",
                "UMP DB без изменений",
                WHERE_MUT,
                "Разработчик",
            ),
            ("9", "Скрипт, содержащий сопутствующие изменения в БД (миграция)", "Нет необходимости", "—", WHERE_MIG, "Разработчик"),
            (
                "10",
                "Заглушки для внешних зависимостей",
                "Пока мок. План: реальный AC в квартале. Документ должен быть создан в prepare; своей папки на НТ нет.",
                "Мок delete на НТ",
                WHERE_STUB,
                "Разработчик",
            ),
            (
                "11",
                "Ожидаемое время отклика замоканных внешних зависимостей",
                "Несколько секунд на document (TBD).",
                "мок delete ~несколько сек × N",
                WHERE_LAT,
                "Разработчик",
            ),
            (
                "12",
                "Сценарий тестирования",
                (
                    "1) Signing timeout (PT25M) → main Call Activity delete (mock) → BANK_REJECT (~5%, из них ~95% signing).\n"
                    "2) Payment timeout (PT5M) → delete (редко).\n"
                    "3) Main PT30M → delete (аварийный).\n"
                    "4) После реального AC: невалидный UUID."
                ),
                "Моделировать ~5% timeout-поток",
                WHERE_SCEN,
                "Аналитик",
            ),
        ],
        "bpmn_section": "Шаги из ump-delete-documents-ncins-pa.bpmn — STATIC vs DYNAMIC",
        "bpmn_rows": [
            ("1", "Start", "Start", "startEvent", "businessKey, productCode", "DYNAMIC start"),
            ("2", "service-task-get-params", "Получить данные", "get-application-data", "processType==DELETE_DOCS", "STATIC processType=DELETE_DOCS; DYNAMIC businessKey"),
            ("3", "service-task-delete-documents", "Удалить из AC", "ump-delete-documents-ncins-pa.delete-documents", "multiInstance sequential over documents; syscode=UMP; out acDocuments", "STATIC syscode; DYNAMIC documents/document"),
            ("4", "End", "End", "endEvent", "—", "—"),
        ],
    }
)


def write_index() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Оглавление"
    ws["A1"] = "НТ пререквизиты — пакет ncins (структура = образец ump-onboarding-pa)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Одинаковые вопросы (пререквизиты 1–12). Ответы — СА + бэкенд + сверка с BPMN PR. TBD где данных нет."
    ws["A2"].font = SUB_FONT
    ws["A3"] = "Файл"
    ws["B3"] = "Процесс"
    style_header_row(ws, 3, 2)
    rows = [
        ("НТ пререквизиты __ Worker ump-main-ma-ncins-pa.xlsx", "Управление процессом мультизаявки"),
        ("НТ пререквизиты __ Worker ump-prepare-documents-ncins-pa.xlsx", "Формирования документов"),
        ("НТ пререквизиты __ Worker ump-signing-documents-ncins-pa.xlsx", "Подписание документов"),
        ("НТ пререквизиты __ Worker ump-payment-ncins-pa.xlsx", "Оплата страховки"),
        ("НТ пререквизиты __ Worker ump-finalisation-ncins-pa.xlsx", "Финализация заявки"),
        ("НТ пререквизиты __ Worker ump-delete-documents-ncins-pa.xlsx", "Удаление документов"),
    ]
    for i, (a, b) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=a).font = CELL_FONT
        ws.cell(row=i, column=2, value=b).font = CELL_FONT
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 40
    wb.save(OUT / "00-README-index.xlsx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_index()
    for meta in PROCESSES:
        wb = build_workbook(meta)
        path = OUT / meta["file"]
        wb.save(path)
        print("wrote", path)


if __name__ == "__main__":
    main()
